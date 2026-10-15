"""Exportación del resultado de una auditoría a los formatos de entrega.

Siempre se escribe resultado.json (fuente completa, incluye cobertura);
los demás formatos se generan solo de los pedidos. Los datos exportados son
exactamente los que el sitio entregó: las secciones sin datos se declaran
como no disponibles con su motivo.
"""
import html as htmllib
import json
import re
from datetime import datetime
from pathlib import Path

FORMATOS = {"xlsx": "Excel (.xlsx)", "pdf": "PDF (.pdf)", "md": "Markdown (.md)",
            "csv": "CSV (uno por tabla)", "html": "HTML (.html)", "json": "JSON (.json)"}


def _slug(titulo):
    return re.sub(r"[^a-z0-9]+", "-", (titulo or "seccion").lower()).strip("-")[:50] or "seccion"


def _lineas_seccion(sec):
    """Representación en líneas de texto de una sección (para MD/HTML/PDF)."""
    L = []
    if sec.get("descripcion"):
        L.append(sec["descripcion"])
    if sec.get("texto"):
        L.append(sec["texto"])
    if sec.get("detalle"):
        L.append(f"Nota: {sec['detalle']}")
    if sec.get("no_disponible"):
        L.append(f"DATO NO DISPONIBLE ({sec['no_disponible']}): "
                 f"{sec.get('detalle') or 'el sitio no entregó esta información'}")
    elif sec.get("tabla"):
        cols = list(sec["tabla"][0].keys())
        L.append("| " + " | ".join(cols) + " |")
        L.append("|" + "---|" * len(cols))
        for fila in sec["tabla"]:
            L.append("| " + " | ".join(str(fila.get(c, "")) for c in cols) + " |")
    return L


def _tabla_cobertura(resultado):
    filas = [{"Módulo": c["modulo"], "Sección": c["seccion"], "Estado": c["estado"],
              "Detalle": c["detalle"]} for c in resultado.get("cobertura", [])]
    return filas


def a_markdown(resultado):
    L = ["# Informe de auditoría de sitio", "",
         f"- **Flujo:** {resultado.get('nombre', '')} ({resultado.get('flujo_id', '')})",
         f"- **Sitio:** {resultado.get('sitio', '')}",
         f"- **Generado:** {resultado.get('generado', '')}",
         f"- **Nivel de acceso:** {resultado.get('nivel_acceso', {}).get('nivel', '')}", ""]
    conn = resultado.get("conexion") or {}
    L += ["## Conexión", "",
          f"- HTTP: {conn.get('http')} — latencia {conn.get('latencia_ms')} ms",
          f"- Título: {conn.get('titulo') or '(sin título)'}",
          f"- Estado: {conn.get('estado')} {conn.get('error') or ''}", ""]
    acc = resultado.get("acceso") or {}
    if acc:
        L += ["## Acceso utilizado", ""]
        for quien, info in acc.items():
            estado = info.get("estado", "")
            msg = info.get("mensaje_sitio") or info.get("mensaje") or info.get("detalle") or ""
            L.append(f"- **{quien}**: {estado}. {msg}".rstrip())
        L.append("")
    for titulo, sec in (resultado.get("secciones") or {}).items():
        L += [f"## {titulo}", ""]
        L += [l for l in _lineas_seccion(sec) if l.strip()]
        L.append("")
    L += ["## Cobertura de la auditoría", "",
          "Qué datos reales se obtuvieron y qué no se pudo obtener, con el motivo devuelto por el sitio.", ""]
    for fila in _tabla_cobertura(resultado):
        L.append(f"- **{fila['Módulo']} / {fila['Sección']}** — {fila['Estado']}: {fila['Detalle']}")
    L.append("")
    return "\n".join(L) + "\n"


def a_html(resultado):
    e = htmllib.escape
    partes = ["<html><head><meta charset='utf-8'><title>Informe de auditoría de sitio</title>",
              "<style>body{font-family:sans-serif;max-width:950px;margin:2em auto}"
              "table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:4px 8px}"
              ".nodisp{color:#a00;font-weight:bold}</style></head><body>",
              "<h1>Informe de auditoría de sitio</h1>",
              f"<p><b>Flujo:</b> {e(str(resultado.get('nombre', '')))} "
              f"({e(str(resultado.get('flujo_id', '')))})<br>"
              f"<b>Sitio:</b> {e(str(resultado.get('sitio', '')))}<br>"
              f"<b>Generado:</b> {e(str(resultado.get('generado', '')))}<br>",
              f"<b>Nivel de acceso:</b> {e(str(resultado.get('nivel_acceso', {}).get('nivel', '')))}</p>"]
    conn = resultado.get("conexion") or {}
    partes.append(f"<h2>Conexión</h2><p>HTTP {e(str(conn.get('http')))} — "
                  f"{e(str(conn.get('latencia_ms')))} ms — {e(str(conn.get('titulo') or ''))} — "
                  f"estado {e(str(conn.get('estado')))} {e(str(conn.get('error') or ''))}</p>")
    for titulo, sec in (resultado.get("secciones") or {}).items():
        partes.append(f"<h2>{e(titulo)}</h2>")
        for ln in _lineas_seccion(sec):
            if ln.startswith("|"):
                partes.append("<p><code>" + e(ln) + "</code></p>")
                continue
            clase = " class='nodisp'" if ln.startswith("DATO NO DISPONIBLE") else ""
            partes.append(f"<p{clase}>{e(ln)}</p>")
        if sec.get("tabla"):
            partes.append("<table><tr>" + "".join(
                f"<th>{e(str(k))}</th>" for k in sec["tabla"][0]) + "</tr>" +
                "".join("<tr>" + "".join(f"<td>{e(str(v))}</td>" for v in fila.values()) + "</tr>"
                        for fila in sec["tabla"]) + "</table>")
    partes.append("<h2>Cobertura de la auditoría</h2><table><tr><th>Módulo</th><th>Sección</th>"
                  "<th>Estado</th><th>Detalle</th></tr>")
    for fila in _tabla_cobertura(resultado):
        partes.append(f"<tr><td>{e(fila['Módulo'])}</td><td>{e(fila['Sección'])}</td>"
                      f"<td>{e(fila['Estado'])}</td><td>{e(fila['Detalle'])}</td></tr>")
    partes.append("</table></body></html>")
    return "\n".join(partes)


def a_pdf(resultado):
    from fpdf import FPDF

    def t(s):  # fpdf con fuentes core solo admite latin-1
        return str(s).encode("latin-1", "replace").decode("latin-1")

    pdf = FPDF()
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, t("Informe de auditoria de sitio"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    pdf.multi_cell(0, 6, t(f"Flujo: {resultado.get('nombre', '')}  ({resultado.get('flujo_id', '')})\n"
                           f"Sitio: {resultado.get('sitio', '')}\n"
                           f"Generado: {resultado.get('generado', '')}\n"
                           f"Nivel de acceso: {resultado.get('nivel_acceso', {}).get('nivel', '')}"),
                   new_x="LMARGIN", new_y="NEXT")

    def seccion(titulo, lineas, tabla=None):
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 7, t(titulo), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=9)
        for ln in lineas:
            pdf.multi_cell(0, 5, t(ln), new_x="LMARGIN", new_y="NEXT")
        if tabla:
            cols = list(tabla[0].keys())
            pdf.set_font("Helvetica", "B", 8)
            ancho = 190 / len(cols)
            for c in cols:
                pdf.cell(ancho, 5, t(c)[:20], border=1)
            pdf.ln()
            pdf.set_font("Helvetica", size=7)
            for fila in tabla[:120]:
                for c in cols:
                    pdf.cell(ancho, 4, t(str(fila.get(c, "")))[:26], border=1)
                pdf.ln()
            if len(tabla) > 120:
                pdf.multi_cell(0, 5, t(f"(se muestran 120 de {len(tabla)} filas; "
                                       f"el listado completo está en resultado.json)"),
                               new_x="LMARGIN", new_y="NEXT")

    for titulo, sec in (resultado.get("secciones") or {}).items():
        seccion(titulo, [l for l in _lineas_seccion(sec) if l.strip() and not l.startswith("|")],
                sec.get("tabla"))
    seccion("Cobertura de la auditoria", ["Qué se obtuvo y qué no, con el motivo real del sitio."],
            _tabla_cobertura(resultado))
    return bytes(pdf.output())


def a_xlsx(resultado, ruta):
    import openpyxl

    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Resumen")
    ws.append(["Campo", "Valor"])
    for c in ws["A1:B1"]:
        c[0].font = c[1].font = Font(bold=True)
    for campo, valor in [("Flujo", resultado.get("nombre")), ("ID", resultado.get("flujo_id")),
                         ("Sitio", resultado.get("sitio")), ("Generado", resultado.get("generado")),
                         ("Nivel de acceso", resultado.get("nivel_acceso", {}).get("nivel")),
                         ("Conexión", f"HTTP {resultado.get('conexion', {}).get('http')} "
                          f"{resultado.get('conexion', {}).get('error') or ''}".strip())]:
        ws.append([campo, valor])
    ws_c = wb.create_sheet("Cobertura")
    filas_cob = _tabla_cobertura(resultado)
    ws_c.append(list(filas_cob[0].keys()) if filas_cob else ["Módulo", "Sección", "Estado", "Detalle"])
    for c in ws_c[1]:
        c.font = Font(bold=True)
    for fila in filas_cob:
        ws_c.append(list(fila.values()))
    for titulo, sec in (resultado.get("secciones") or {}).items():
        hoja = _slug(titulo)[:28]
        ws_s = wb.create_sheet(hoja)
        if sec.get("tabla"):
            cols = list(sec["tabla"][0].keys())
            ws_s.append(cols)
            for c in ws_s[1]:
                c.font = Font(bold=True)
            for fila in sec["tabla"]:
                ws_s.append([fila.get(c) for c in cols])
        lineas = [l for l in _lineas_seccion(sec) if l.strip() and not l.startswith("|")]
        if lineas:
            if sec.get("tabla"):
                ws_s.append([])
            ws_s.append(("Descripción", lineas[0]))
            ws_s.append(("Nota", " ".join(lineas[1:]) or "-"))
    wb.save(ruta)


def exportar_resultado(resultado, carpeta, formatos):
    """Escribe los informes en la carpeta dada. Devuelve la lista de archivos."""
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    generados = []
    ruta_json = carpeta / "resultado.json"
    ruta_json.write_text(json.dumps(resultado, ensure_ascii=False, indent=2), encoding="utf-8")
    generados.append({"archivo": "resultado.json", "ruta": str(ruta_json), "formato": "json"})
    for f in formatos:
        try:
            if f == "json":
                continue
            if f == "md":
                ruta = carpeta / "informe.md"
                ruta.write_text(a_markdown(resultado), encoding="utf-8")
            elif f == "html":
                ruta = carpeta / "informe.html"
                ruta.write_text(a_html(resultado), encoding="utf-8")
            elif f == "pdf":
                ruta = carpeta / "informe.pdf"
                ruta.write_bytes(a_pdf(resultado))
            elif f == "xlsx":
                ruta = carpeta / "informe.xlsx"
                a_xlsx(resultado, ruta)
            elif f == "csv":
                tablas = [(t, s["tabla"]) for t, s in (resultado.get("secciones") or {}).items()
                          if s.get("tabla")]
                if not tablas:
                    continue
                for titulo, tabla in tablas:
                    ruta = carpeta / f"{_slug(titulo)}.csv"
                    import csv

                    with open(ruta, "w", newline="", encoding="utf-8-sig") as fh:
                        w = csv.DictWriter(fh, fieldnames=list(tabla[0].keys()))
                        w.writeheader()
                        w.writerows(tabla)
                    generados.append({"archivo": ruta.name, "ruta": str(ruta), "formato": "csv"})
                continue
            generados.append({"archivo": ruta.name, "ruta": str(ruta), "formato": f})
        except Exception as exc:  # un formato no puede tumbar el informe completo
            resultado.setdefault("errores_exportacion", []).append(
                f"{f}: {type(exc).__name__}: {exc}")
    return generados
